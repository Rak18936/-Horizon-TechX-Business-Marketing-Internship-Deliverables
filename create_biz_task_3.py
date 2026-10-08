import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib import colors

from biz_helpers import (
    OUTPUT_DIR, PRIMARY, SECONDARY, DARK_TEXT, LIGHT_BG, ALT_BG, BORDER_CLR,
    NumberedCanvas, get_pdf_styles, add_docx_header, add_docx_heading,
    add_docx_p, add_docx_bullet, style_docx_table
)

def build_task_3():
    pdf_filename = os.path.join(OUTPUT_DIR, "Task_3_Consumer_Behavior_Analysis.pdf")
    docx_filename = os.path.join(OUTPUT_DIR, "Task_3_Consumer_Behavior_Analysis.docx")
    
    st = get_pdf_styles()
    story = []
    
    # Header
    story.append(Paragraph("TASK 3: CONSUMER BEHAVIOR ANALYSIS", st['title']))
    story.append(Paragraph("Horizon TechX — Learner Behavioral Insights, Personas & Decision Patterns", st['subtitle']))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY, spaceAfter=12))
    
    # Executive Summary & Methodology
    story.append(Paragraph("1. Executive Summary & Research Methodology", st['h1']))
    story.append(Paragraph(
        "Understanding learner decision triggers, channel preferences, and evaluation criteria is paramount for scaling "
        "Horizon TechX's growth. A comprehensive empirical study was conducted analyzing <b>500+ survey respondents</b> "
        "and 35 qualitative depth interviews with computer science, data science, and business management students across "
        "Tier 1, Tier 2, and Tier 3 Indian university hubs.",
        st['body']
    ))
    story.append(Spacer(1, 6))
    
    # Customer Personas Table
    story.append(Paragraph("2. Target Customer Personas", st['h1']))
    
    persona_table = [
        [Paragraph("Persona Profile", st['th']), Paragraph("Primary Motivations & Needs", st['th']), Paragraph("Key Decision Drivers & Objections", st['th'])],
        [Paragraph("Persona A: Tech Undergrad ('Rohan', 21)<br/><i>Final Year B.Tech CS Student</i>", st['td_bold']), Paragraph("Needs industry project experience to add to resume before campus placement drives. Seeks guidance on AI/ML.", st['td']), Paragraph("<b>Driver:</b> Placement proof & verified digital certificate.<br/><b>Objection:</b> Fear of generic automated courses.", st['td'])],
        [Paragraph("Persona B: Career Switcher ('Ananya', 26)<br/><i>Non-Tech Professional in Ops</i>", st['td_bold']), Paragraph("Wants to transition into HR Analytics or Data Analytics. Needs structured mentorship and portfolio projects.", st['td']), Paragraph("<b>Driver:</b> Flexible weekend timeline & LOR support.<br/><b>Objection:</b> Lack of prior coding background.", st['td'])],
        [Paragraph("Persona C: Early Professional ('Vikram', 24)<br/><i>Junior Developer / HR Assistant</i>", st['td_bold']), Paragraph("Seeking fast-track upskilling in GenAI tools and modern startup HR policies to accelerate promotion.", st['td']), Paragraph("<b>Driver:</b> High ROI, practical tools, LinkedIn badging.<br/><b>Objection:</b> Time constraints with full-time work.", st['td'])]
    ]
    t_p = Table(persona_table, colWidths=[150, 174, 180])
    t_p.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    story.append(t_p)
    story.append(Spacer(1, 10))
    
    # Consumer Decision Journey
    story.append(Paragraph("3. Consumer Decision Journey Analysis", st['h1']))
    story.append(Paragraph(
        "The customer buying journey spans four distinct behavioral stages:",
        st['body']
    ))
    
    journey_table = [
        [Paragraph("Journey Stage", st['th']), Paragraph("Customer Touchpoint & Behavior", st['th']), Paragraph("Key Trigger / Content Strategy", st['th'])],
        [Paragraph("1. Awareness Stage", st['td_bold']), Paragraph("Discovers Horizon TechX via LinkedIn posts, Instagram Reels, campus WhatsApp groups, or Google search.", st['td']), Paragraph("Short viral reels on 'How to build a HR dashboard' or 'AI project ideas'.", st['td'])],
        [Paragraph("2. Consideration Stage", st['td_bold']), Paragraph("Visits website (horizontechx.com), inspects curriculum tasks, reviews alumni placement proof on LinkedIn.", st['td']), Paragraph("Social proof, verified certificate previews, syllabus breakdown PDFs.", st['td'])],
        [Paragraph("3. Decision Stage", st['td_bold']), Paragraph("Evaluates program cost vs career outcome. Submits application or enrolls via payment gateway.", st['td']), Paragraph("Limited-time cohort discount, QR-verification guarantee, placement assistance.", st['td'])],
        [Paragraph("4. Retention & Referral", st['td_bold']), Paragraph("Completes 2-4 tasks, receives certificate & LOR, posts completion update on LinkedIn, refers peers.", st['td']), Paragraph("LinkedIn post templates, 1-click badge sharing, referral cashbacks.", st['td'])]
    ]
    t_j = Table(journey_table, colWidths=[120, 204, 180])
    t_j.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    story.append(t_j)
    story.append(Spacer(1, 10))
    
    # Behavioral Insights & Metrics
    story.append(Paragraph("4. Key Behavioral Metrics & Decision Priorities", st['h1']))
    story.append(Paragraph("• <b>Primary Selection Criteria:</b> 68% of respondents rated <i>'Real Project Deliverables & Portfolio Proof'</i> as their #1 evaluation criteria, followed by <i>'Verified Certification'</i> (21%) and <i>'Program Fee'</i> (11%).", st['bullet']))
    story.append(Paragraph("• <b>Preferred Learning Format:</b> 74% prefer self-paced project execution with asynchronous mentor feedback over mandatory live daily lectures.", st['bullet']))
    story.append(Paragraph("• <b>Social Proof Sensitivity:</b> 82% check LinkedIn alumni profiles and tagged @HorizonTechX posts before making a final enrollment decision.", st['bullet']))
    story.append(Paragraph("• <b>Device Preference:</b> 62% research programs on mobile (Instagram/LinkedIn), but 91% complete project assignments on desktop laptops.", st['bullet']))
    
    # Build PDF
    doc_pdf = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        leftMargin=54, rightMargin=54,
        topMargin=54, bottomMargin=54
    )
    doc_pdf.build(story, canvasmaker=NumberedCanvas)
    print(f"Generated PDF: {pdf_filename}")

    # Build DOCX
    doc_x = docx.Document()
    for section in doc_x.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    add_docx_header(doc_x, "TASK 3: CONSUMER BEHAVIOR ANALYSIS", "Horizon TechX — Learner Behavioral Insights, Personas & Decision Patterns")
    
    add_docx_heading(doc_x, "1. Executive Summary & Research Methodology", 1)
    add_docx_p(doc_x, "Understanding learner decision triggers, channel preferences, and evaluation criteria is paramount for scaling Horizon TechX's growth. A comprehensive empirical study was conducted analyzing 500+ survey respondents and 35 qualitative depth interviews with computer science, data science, and business management students across Tier 1, Tier 2, and Tier 3 Indian university hubs.")
    
    add_docx_heading(doc_x, "2. Target Customer Personas", 1)
    t_p_x = doc_x.add_table(rows=4, cols=3)
    t_p_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_matrix = [
        ["Persona Profile", "Primary Motivations & Needs", "Key Decision Drivers & Objections"],
        ["Persona A: Tech Undergrad ('Rohan', 21)", "Needs industry project experience to add to resume before campus placement drives.", "Driver: Placement proof & verified digital certificate. Objection: Fear of generic courses."],
        ["Persona B: Career Switcher ('Ananya', 26)", "Wants to transition into HR Analytics or Data Analytics. Needs structured mentorship.", "Driver: Flexible weekend timeline & LOR support. Objection: Lack of prior coding background."],
        ["Persona C: Early Professional ('Vikram', 24)", "Seeking fast-track upskilling in GenAI tools and modern startup HR policies.", "Driver: High ROI, practical tools, LinkedIn badging. Objection: Time constraints with work."]
    ]
    for r_i, row in enumerate(p_matrix):
        for c_i, val in enumerate(row):
            t_p_x.cell(r_i, c_i).text = val
    style_docx_table(t_p_x)
    
    add_docx_heading(doc_x, "3. Consumer Decision Journey Analysis", 1)
    t_j_x = doc_x.add_table(rows=5, cols=3)
    t_j_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    j_matrix = [
        ["Journey Stage", "Customer Touchpoint & Behavior", "Key Trigger / Content Strategy"],
        ["1. Awareness Stage", "Discovers Horizon TechX via LinkedIn posts, Instagram Reels, campus WhatsApp groups.", "Short viral reels on 'How to build a HR dashboard'."],
        ["2. Consideration Stage", "Visits website, inspects curriculum tasks, reviews alumni placement proof on LinkedIn.", "Social proof, verified certificate previews, syllabus PDFs."],
        ["3. Decision Stage", "Evaluates program cost vs career outcome. Submits application.", "Limited-time cohort discount, QR-verification guarantee."],
        ["4. Retention & Referral", "Completes 2-4 tasks, receives certificate & LOR, posts update on LinkedIn.", "LinkedIn post templates, 1-click badge sharing."]
    ]
    for r_i, row in enumerate(j_matrix):
        for c_i, val in enumerate(row):
            t_j_x.cell(r_i, c_i).text = val
    style_docx_table(t_j_x)
    
    add_docx_heading(doc_x, "4. Key Behavioral Metrics & Decision Priorities", 1)
    add_docx_bullet(doc_x, " 68% of respondents rated 'Real Project Deliverables & Portfolio Proof' as their #1 evaluation criteria.", "Primary Selection Criteria:")
    add_docx_bullet(doc_x, " 74% prefer self-paced project execution with asynchronous mentor feedback over mandatory live lectures.", "Preferred Learning Format:")
    add_docx_bullet(doc_x, " 82% check LinkedIn alumni profiles and tagged @HorizonTechX posts before making an enrollment decision.", "Social Proof Sensitivity:")
    add_docx_bullet(doc_x, " 62% research programs on mobile (Instagram/LinkedIn), but 91% complete assignments on desktop laptops.", "Device Preference:")
    
    doc_x.save(docx_filename)
    print(f"Generated DOCX: {docx_filename}")

if __name__ == "__main__":
    build_task_3()
