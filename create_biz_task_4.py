import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib import colors

from biz_helpers import (
    OUTPUT_DIR, PRIMARY, SECONDARY, DARK_TEXT, LIGHT_BG, ALT_BG, BORDER_CLR,
    NumberedCanvas, get_pdf_styles, add_docx_header, add_docx_heading,
    add_docx_p, add_docx_bullet, style_docx_table
)

def build_task_4():
    pdf_filename = os.path.join(OUTPUT_DIR, "Task_4_Business_Growth_Campaign_Plan.pdf")
    docx_filename = os.path.join(OUTPUT_DIR, "Task_4_Business_Growth_Campaign_Plan.docx")
    
    st = get_pdf_styles()
    story = []
    
    # Header
    story.append(Paragraph("TASK 4: BUSINESS GROWTH CAMPAIGN PLAN", st['title']))
    story.append(Paragraph("Horizon TechX — 90-Day 'TechX Accelerate 2026' Growth Campaign Strategy", st['subtitle']))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY, spaceAfter=12))
    
    # Executive Campaign Summary
    story.append(Paragraph("1. Campaign Overview & Core Objectives", st['h1']))
    story.append(Paragraph(
        "To rapidly expand market presence, drive user acquisition, and establish Horizon TechX as the #1 practical project internship "
        "platform in India and Southeast Asia, we present the 90-Day <b>'TechX Accelerate 2026' Growth Campaign</b>. "
        "The campaign operates across top-of-funnel brand awareness, middle-of-funnel lead nurturing, and bottom-of-funnel enrollment conversion.",
        st['body']
    ))
    story.append(Spacer(1, 6))
    
    obj_table = [
        [Paragraph("Campaign Objective", st['th']), Paragraph("Target Benchmark", st['th']), Paragraph("Measurement Metric / KPI", st['th'])],
        [Paragraph("1. Brand Awareness & Impressions", st['td_bold']), Paragraph("2.5 Million Total Impressions", st['td']), Paragraph("Ad Impressions, Video Views (>75%), Social Reach.", st['td'])],
        [Paragraph("2. Lead Generation (Signups)", st['td_bold']), Paragraph("15,000 Verified Leads", st['td']), Paragraph("Form submissions, eBook downloads, Webinar RSVPs.", st['td'])],
        [Paragraph("3. Paid Cohort Conversions", st['td_bold']), Paragraph("2,200 Program Enrollees", st['td']), Paragraph("Paid enrollments, Conversion Rate (>14.6%).", st['td'])],
        [Paragraph("4. Customer Acquisition Cost (CAC)", st['td_bold']), Paragraph("< ₹115 per acquisition", st['td']), Paragraph("Blended CAC across Paid + Organic channels.", st['td'])],
        [Paragraph("5. Return on Ad Spend (ROAS)", st['td_bold']), Paragraph("3.8x Projected ROAS", st['td']), Paragraph("Total Revenue generated / Total Marketing Spend.", st['td'])]
    ]
    t_o = Table(obj_table, colWidths=[150, 154, 200])
    t_o.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE')
    ]))
    story.append(t_o)
    story.append(Spacer(1, 10))
    
    # Audience Targeting & Media Mix
    story.append(Paragraph("2. Audience Targeting & Channel Execution Strategy", st['h1']))
    story.append(Paragraph(
        "The campaign targets three primary demographics using tailored message angles and media assets:",
        st['body']
    ))
    
    story.append(Paragraph("• <b>Audience Segment A (Tech & CS Undergrads):</b> Ages 18-23, interest in Python, AI, React, Data Science, internship opportunities. <i>Channel: Instagram Ads & Campus WhatsApp Groups.</i>", st['bullet']))
    story.append(Paragraph("• <b>Audience Segment B (Management & HR Aspirants):</b> Ages 20-25, interest in HR Management, POSH compliance, People Analytics, MBA. <i>Channel: LinkedIn Sponsored Content & Search Ads.</i>", st['bullet']))
    story.append(Paragraph("• <b>Audience Segment C (Early Career Pros):</b> Ages 22-28, seeking career pivots into AI or Data Analytics. <i>Channel: Google Search Ads & YouTube Tech Tutorials.</i>", st['bullet']))
    
    story.append(Spacer(1, 8))
    
    # Budget Allocation & Financial Projections
    story.append(Paragraph("3. Campaign Budget Allocation (Total: ₹2,50,000)", st['h1']))
    
    budget_table = [
        [Paragraph("Marketing Channel / Expense", st['th']), Paragraph("Budget (INR)", st['th']), Paragraph("Share (%)", st['th']), Paragraph("Expected Deliverables & Outcomes", st['th'])],
        [Paragraph("Meta Paid Ads (Instagram/Facebook)", st['td_bold']), Paragraph("₹1,00,000", st['td']), Paragraph("40%", st['td']), Paragraph("1.2M Impressions, 8,000 Leads (CPL ~₹12.50).", st['td'])],
        [Paragraph("LinkedIn Performance Ads", st['td_bold']), Paragraph("₹65,000", st['td']), Paragraph("26%", st['td']), Paragraph("300k B2B/HR Impressions, 2,500 High-Intent Leads.", st['td'])],
        [Paragraph("Google Search (PPC)", st['td_bold']), Paragraph("₹40,000", st['td']), Paragraph("16%", st['td']), Paragraph("High-intent searches ('best HR internship online').", st['td'])],
        [Paragraph("Campus Webinars & Influencers", st['td_bold']), Paragraph("₹30,000", st['td']), Paragraph("12%", st['td']), Paragraph("10 Creator videos & 5 university masterclasses.", st['td'])],
        [Paragraph("Creative Production & Software", st['td_bold']), Paragraph("₹15,000", st['td']), Paragraph("6%", st['td']), Paragraph("Ad design templates, video editing, tracking tools.", st['td'])]
    ]
    t_b = Table(budget_table, colWidths=[150, 80, 54, 220])
    t_b.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE')
    ]))
    story.append(t_b)
    story.append(Spacer(1, 10))
    
    # 90-Day Execution Roadmap
    story.append(Paragraph("4. 90-Day Phased Execution Roadmap", st['h1']))
    
    roadmap_table = [
        [Paragraph("Phase & Timeline", st['th']), Paragraph("Core Execution Activities", st['th']), Paragraph("Milestone Output", st['th'])],
        [Paragraph("Phase 1: Days 1–15<br/><i>Setup & Creative Testing</i>", st['td_bold']), Paragraph("Finalize landing pages, install pixel tracking, launch A/B tests for 10 ad creatives and headline variations.", st['td']), Paragraph("Identify top 3 winning ad creatives.", st['td'])],
        [Paragraph("Phase 2: Days 16–60<br/><i>Scale & Optimization</i>", st['td_bold']), Paragraph("Scale winning ad sets, host weekly live masterclasses, activate college campus ambassador referral drives.", st['td']), Paragraph("Achieve 10,000 total signups.", st['td'])],
        [Paragraph("Phase 3: Days 61–90<br/><i>Retargeting & Conversion</i>", st['td_bold']), Paragraph("Deploy dynamic retargeting ads to un-enrolled leads, send urgency emails for cohort deadline, collect testimonials.", st['td']), Paragraph("Hit 2,200 paid conversions (3.8x ROAS).", st['td'])]
    ]
    t_r = Table(roadmap_table, colWidths=[120, 244, 140])
    t_r.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    story.append(t_r)
    
    # Build PDF
    doc_pdf = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        leftMargin=54, rightMargin=54,
        topMargin=54, bottomMargin=54
    )
    doc_pdf.build(story, canvasmaker=NumberedCanvas)
    print(f"Generated PDF: {pdf_filename}")

    # Build DOCX
    doc_x = docx.Document()
    for section in doc_x.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    add_docx_header(doc_x, "TASK 4: BUSINESS GROWTH CAMPAIGN PLAN", "Horizon TechX — 90-Day 'TechX Accelerate 2026' Growth Campaign Strategy")
    
    add_docx_heading(doc_x, "1. Campaign Overview & Core Objectives", 1)
    add_docx_p(doc_x, "To rapidly expand market presence, drive user acquisition, and establish Horizon TechX as the #1 practical project internship platform in India and Southeast Asia, we present the 90-Day 'TechX Accelerate 2026' Growth Campaign. The campaign operates across top-of-funnel brand awareness, middle-of-funnel lead nurturing, and bottom-of-funnel enrollment conversion.")
    
    t_o_x = doc_x.add_table(rows=6, cols=3)
    t_o_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    o_matrix = [
        ["Campaign Objective", "Target Benchmark", "Measurement Metric / KPI"],
        ["1. Brand Awareness & Impressions", "2.5 Million Total Impressions", "Ad Impressions, Video Views (>75%), Social Reach."],
        ["2. Lead Generation (Signups)", "15,000 Verified Leads", "Form submissions, eBook downloads, Webinar RSVPs."],
        ["3. Paid Cohort Conversions", "2,200 Program Enrollees", "Paid enrollments, Conversion Rate (>14.6%)."],
        ["4. Customer Acquisition Cost (CAC)", "< ₹115 per acquisition", "Blended CAC across Paid + Organic channels."],
        ["5. Return on Ad Spend (ROAS)", "3.8x Projected ROAS", "Total Revenue generated / Total Marketing Spend."]
    ]
    for r_i, row in enumerate(o_matrix):
        for c_i, val in enumerate(row):
            t_o_x.cell(r_i, c_i).text = val
    style_docx_table(t_o_x)
    
    add_docx_heading(doc_x, "2. Audience Targeting & Channel Execution Strategy", 1)
    add_docx_bullet(doc_x, " Ages 18-23, interest in Python, AI, React, Data Science, internship opportunities. Channel: Instagram Ads & Campus WhatsApp Groups.", "Audience Segment A (Tech Undergrads):")
    add_docx_bullet(doc_x, " Ages 20-25, interest in HR Management, POSH compliance, People Analytics, MBA. Channel: LinkedIn Ads & Search.", "Audience Segment B (HR Aspirants):")
    add_docx_bullet(doc_x, " Ages 22-28, seeking career pivots into AI or Data Analytics. Channel: Google Search Ads & YouTube Tutorials.", "Audience Segment C (Early Career Pros):")
    
    add_docx_heading(doc_x, "3. Campaign Budget Allocation (Total: ₹2,50,000)", 1)
    t_b_x = doc_x.add_table(rows=6, cols=4)
    t_b_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_matrix = [
        ["Marketing Channel / Expense", "Budget (INR)", "Share (%)", "Expected Deliverables & Outcomes"],
        ["Meta Paid Ads (Instagram/Facebook)", "₹1,00,000", "40%", "1.2M Impressions, 8,000 Leads (CPL ~₹12.50)."],
        ["LinkedIn Performance Ads", "₹65,000", "26%", "300k B2B/HR Impressions, 2,500 High-Intent Leads."],
        ["Google Search (PPC)", "₹40,000", "16%", "High-intent searches ('best HR internship online')."],
        ["Campus Webinars & Influencers", "₹30,000", "12%", "10 Creator videos & 5 university masterclasses."],
        ["Creative Production & Software", "₹15,000", "6%", "Ad design templates, video editing, tracking tools."]
    ]
    for r_i, row in enumerate(b_matrix):
        for c_i, val in enumerate(row):
            t_b_x.cell(r_i, c_i).text = val
    style_docx_table(t_b_x)
    
    add_docx_heading(doc_x, "4. 90-Day Phased Execution Roadmap", 1)
    t_r_x = doc_x.add_table(rows=4, cols=3)
    t_r_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_matrix = [
        ["Phase & Timeline", "Core Execution Activities", "Milestone Output"],
        ["Phase 1: Days 1–15 (Setup)", "Finalize landing pages, install pixel tracking, launch A/B tests for 10 ad creatives.", "Identify top 3 winning ad creatives."],
        ["Phase 2: Days 16–60 (Scale)", "Scale winning ad sets, host weekly live masterclasses, activate campus ambassador drives.", "Achieve 10,000 total signups."],
        ["Phase 3: Days 61–90 (Retargeting)", "Deploy dynamic retargeting ads to un-enrolled leads, send urgency emails for cohort deadline.", "Hit 2,200 paid conversions (3.8x ROAS)."]
    ]
    for r_i, row in enumerate(r_matrix):
        for c_i, val in enumerate(row):
            t_r_x.cell(r_i, c_i).text = val
    style_docx_table(t_r_x)
    
    doc_x.save(docx_filename)
    print(f"Generated DOCX: {docx_filename}")

if __name__ == "__main__":
    build_task_4()
