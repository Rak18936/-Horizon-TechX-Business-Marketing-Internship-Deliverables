import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib import colors

from biz_helpers import (
    OUTPUT_DIR, PRIMARY, SECONDARY, DARK_TEXT, LIGHT_BG, ALT_BG, BORDER_CLR,
    NumberedCanvas, get_pdf_styles, add_docx_header, add_docx_heading,
    add_docx_p, add_docx_bullet, style_docx_table
)

def build_task_2():
    pdf_filename = os.path.join(OUTPUT_DIR, "Task_2_Brand_Positioning_and_Growth_Strategy.pdf")
    docx_filename = os.path.join(OUTPUT_DIR, "Task_2_Brand_Positioning_and_Growth_Strategy.docx")
    
    st = get_pdf_styles()
    story = []
    
    # Header
    story.append(Paragraph("TASK 2: BRAND POSITIONING & GROWTH STRATEGY", st['title']))
    story.append(Paragraph("Horizon TechX — Brand Positioning Framework & Multi-Channel Engagement Blueprint", st['subtitle']))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY, spaceAfter=12))
    
    # Brand Purpose & Positioning
    story.append(Paragraph("1. Core Brand Positioning Framework", st['h1']))
    story.append(Paragraph(
        "<b>Brand Purpose:</b> To bridge the critical gap between academic education and industry performance by providing "
        "ambitious learners with practical, project-based internship experience, verified credentials, and career readiness support.",
        st['body']
    ))
    story.append(Paragraph(
        "<b>Brand Positioning Statement:</b> <i>'For college students and early-career tech professionals who seek real-world work experience, "
        "Horizon TechX is the premier practical skill acceleration platform that delivers verifiable project portfolios and placement support, "
        "unlike traditional video-only courses that lack practical application.'</i>",
        st['callout']
    ))
    story.append(Spacer(1, 8))
    
    # Value Proposition Pillars
    story.append(Paragraph("2. Unique Selling Proposition (USP) & Value Pillars", st['h1']))
    
    brand_pillars = [
        [Paragraph("Brand Pillar", st['th']), Paragraph("Core Promise", st['th']), Paragraph("Proof Points & Features", st['th'])],
        [Paragraph("1. Practical Project-First Learning", st['td_bold']), Paragraph("Zero fluff. 100% real-world assignments.", st['td']), Paragraph("Hands-on tasks: HR policies, AI research papers, live dashboards, coding repos.", st['td'])],
        [Paragraph("2. Verifiable Credentialing", st['td_bold']), Paragraph("Tamper-proof, employer-ready validation.", st['td']), Paragraph("QR-verified certificates, Unique Student IDs, LinkedIn shareable digital badges.", st['td'])],
        [Paragraph("3. Fast-Track Career Growth", st['td_bold']), Paragraph("Direct transition into high-demand roles.", st['td']), Paragraph("Resume building support, Letter of Recommendation (LOR), placement assistance.", st['td'])],
        [Paragraph("4. Accessible & Inclusive", st['td_bold']), Paragraph("Democratizing high-quality tech training.", st['td']), Paragraph("Affordable pricing, self-paced flexible tracks, remote hybrid participation.", st['td'])]
    ]
    t_bp = Table(brand_pillars, colWidths=[130, 150, 224])
    t_bp.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    story.append(t_bp)
    story.append(Spacer(1, 10))
    
    # Multi-Channel Engagement Strategy
    story.append(Paragraph("3. Multi-Channel Customer Engagement Strategy", st['h1']))
    story.append(Paragraph(
        "To build strong brand affinity and sustain organic customer acquisition, Horizon TechX deploys a multi-channel engagement model:",
        st['body']
    ))
    
    story.append(Paragraph("• <b>LinkedIn B2B & Thought Leadership (Primary Channel):</b> Publish intern success stories, project showcases, industry research insights, and tag corporate partners to drive viral organic reach.", st['bullet']))
    story.append(Paragraph("• <b>Instagram & Community Engagement (Youth Appeal):</b> Share short-form video explainers (Reels), intern spot-checks, behind-the-scenes culture, and host Q&A sessions on career guidance.", st['bullet']))
    story.append(Paragraph("• <b>YouTube Content Hub (Educational Value):</b> Host deep-dive tutorial series, project walkthroughs, and alumni interview series to capture organic search traffic.", st['bullet']))
    story.append(Paragraph("• <b>Email Nurture Sequences (Lifecycle Retention):</b> Automated drip campaigns delivering project tips, weekly learning recaps, and exclusive job placement alerts.", st['bullet']))
    
    story.append(Spacer(1, 8))
    
    # Growth Strategy Levers
    story.append(Paragraph("4. Growth Strategy & Scalable Acquisition Levers", st['h1']))
    
    growth_levers = [
        [Paragraph("Growth Lever", st['th']), Paragraph("Strategy & Execution Mechanism", st['th']), Paragraph("Expected Impact / Target KPI", st['th'])],
        [Paragraph("Referral Loop Program", st['td_bold']), Paragraph("Provide existing interns with ₹500 referral bonus or learning credits for every peer enrolled.", st['td']), Paragraph("30% increase in word-of-mouth acquisition.", st['td'])],
        [Paragraph("Campus Ambassador Chapter", st['td_bold']), Paragraph("Appoint 200+ student leaders across top tech universities to host campus hackathons.", st['td']), Paragraph("5,000+ new organic leads monthly.", st['td'])],
        [Paragraph("SEO Content Engine", st['td_bold']), Paragraph("Publish long-tail keyword guides ('How to get HR internship', 'AI project portfolio guide').", st['td']), Paragraph("100,000+ monthly organic blog visits.", st['td'])],
        [Paragraph("Performance Marketing", st['td_bold']), Paragraph("Targeted Meta & LinkedIn ads focusing on outcome-oriented messaging and alumni proof.", st['td']), Paragraph("Maintain Cost Per Lead (CPL) under ₹120.", st['td'])]
    ]
    t_gl = Table(growth_levers, colWidths=[120, 244, 140])
    t_gl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    story.append(t_gl)
    
    # Build PDF
    doc_pdf = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        leftMargin=54, rightMargin=54,
        topMargin=54, bottomMargin=54
    )
    doc_pdf.build(story, canvasmaker=NumberedCanvas)
    print(f"Generated PDF: {pdf_filename}")

    # Build DOCX
    doc_x = docx.Document()
    for section in doc_x.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    add_docx_header(doc_x, "TASK 2: BRAND POSITIONING & GROWTH STRATEGY", "Horizon TechX — Brand Positioning Framework & Multi-Channel Engagement Blueprint")
    
    add_docx_heading(doc_x, "1. Core Brand Positioning Framework", 1)
    add_docx_p(doc_x, "Brand Purpose: To bridge the critical gap between academic education and industry performance by providing ambitious learners with practical, project-based internship experience, verified credentials, and career readiness support.")
    add_docx_p(doc_x, "Brand Positioning Statement: 'For college students and early-career tech professionals who seek real-world work experience, Horizon TechX is the premier practical skill acceleration platform that delivers verifiable project portfolios and placement support, unlike traditional video-only courses that lack practical application.'")
    
    add_docx_heading(doc_x, "2. Unique Selling Proposition (USP) & Value Pillars", 1)
    t_bp_x = doc_x.add_table(rows=5, cols=3)
    t_bp_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    bp_matrix = [
        ["Brand Pillar", "Core Promise", "Proof Points & Features"],
        ["1. Practical Project-First Learning", "Zero fluff. 100% real-world assignments.", "Hands-on tasks: HR policies, AI research papers, live dashboards, coding repos."],
        ["2. Verifiable Credentialing", "Tamper-proof, employer-ready validation.", "QR-verified certificates, Unique Student IDs, LinkedIn shareable digital badges."],
        ["3. Fast-Track Career Growth", "Direct transition into high-demand roles.", "Resume building support, Letter of Recommendation (LOR), placement assistance."],
        ["4. Accessible & Inclusive", "Democratizing high-quality tech training.", "Affordable pricing, self-paced flexible tracks, remote hybrid participation."]
    ]
    for r_i, row in enumerate(bp_matrix):
        for c_i, val in enumerate(row):
            t_bp_x.cell(r_i, c_i).text = val
    style_docx_table(t_bp_x)
    
    add_docx_heading(doc_x, "3. Multi-Channel Customer Engagement Strategy", 1)
    add_docx_bullet(doc_x, " Publish intern success stories, project showcases, industry research insights, and tag corporate partners to drive viral organic reach.", "LinkedIn B2B & Thought Leadership:")
    add_docx_bullet(doc_x, " Share short-form video explainers (Reels), intern spot-checks, behind-the-scenes culture, and host Q&A sessions on career guidance.", "Instagram & Community Engagement:")
    add_docx_bullet(doc_x, " Host deep-dive tutorial series, project walkthroughs, and alumni interview series to capture organic search traffic.", "YouTube Content Hub:")
    add_docx_bullet(doc_x, " Automated drip campaigns delivering project tips, weekly learning recaps, and exclusive job placement alerts.", "Email Nurture Sequences:")
    
    add_docx_heading(doc_x, "4. Growth Strategy & Scalable Acquisition Levers", 1)
    t_gl_x = doc_x.add_table(rows=5, cols=3)
    t_gl_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    gl_matrix = [
        ["Growth Lever", "Strategy & Execution Mechanism", "Expected Impact / Target KPI"],
        ["Referral Loop Program", "Provide existing interns with ₹500 referral bonus or learning credits for every peer enrolled.", "30% increase in word-of-mouth acquisition."],
        ["Campus Ambassador Chapter", "Appoint 200+ student leaders across top tech universities to host campus hackathons.", "5,000+ new organic leads monthly."],
        ["SEO Content Engine", "Publish long-tail keyword guides ('How to get HR internship', 'AI project portfolio guide').", "100,000+ monthly organic blog visits."],
        ["Performance Marketing", "Targeted Meta & LinkedIn ads focusing on outcome-oriented messaging and alumni proof.", "Maintain Cost Per Lead (CPL) under ₹120."]
    ]
    for r_i, row in enumerate(gl_matrix):
        for c_i, val in enumerate(row):
            t_gl_x.cell(r_i, c_i).text = val
    style_docx_table(t_gl_x)
    
    doc_x.save(docx_filename)
    print(f"Generated DOCX: {docx_filename}")

if __name__ == "__main__":
    build_task_2()
