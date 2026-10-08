import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas

OUTPUT_DIR = r"C:\Users\laks1\HorizonTechX_Business_Growth_Tasks"
os.makedirs(OUTPUT_DIR, exist_ok=True)

# Brand Color Palette: Premium Indigo & Violet Theme
PRIMARY = colors.HexColor("#4338CA")     # Indigo primary
SECONDARY = colors.HexColor("#0D9488")   # Teal accent
DARK_TEXT = colors.HexColor("#1E293B")   # Dark Slate
LIGHT_BG = colors.HexColor("#F8FAFC")    # Soft slate background
ALT_BG = colors.HexColor("#F1F5F9")      # Table alt row
BORDER_CLR = colors.HexColor("#CBD5E1")  # Light gray border
ACCENT_MUTED = colors.HexColor("#EEF2FF")

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super().showPage()
        super().save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header (pages > 1)
        if self._pageNumber > 1:
            self.drawString(54, 750, "Horizon TechX — Business Growth & Market Research Deliverables")
            self.setStrokeColor(colors.HexColor("#E2E8F0"))
            self.setLineWidth(0.5)
            self.line(54, 742, 558, 742)

        # Footer
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 36, page_text)
        self.drawString(54, 36, "Horizon TechX Internship Deliverable | Confidential")
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(54, 48, 558, 48)
        self.restoreState()


def get_pdf_styles():
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=PRIMARY,
        spaceAfter=6
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=SECONDARY,
        spaceAfter=14
    )

    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=PRIMARY,
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10.5,
        leading=14,
        textColor=SECONDARY,
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=DARK_TEXT,
        spaceAfter=5
    )

    bullet_style = ParagraphStyle(
        'Bullet_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=DARK_TEXT,
        leftIndent=12,
        spaceAfter=3
    )

    callout_style = ParagraphStyle(
        'Callout',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=13,
        textColor=PRIMARY,
        backColor=ACCENT_MUTED,
        borderColor=SECONDARY,
        borderWidth=0.8,
        borderPadding=6,
        spaceBefore=6,
        spaceAfter=6
    )

    table_header = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )

    table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=DARK_TEXT
    )

    table_cell_bold = ParagraphStyle(
        'TableCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        textColor=DARK_TEXT
    )

    return {
        'title': title_style,
        'subtitle': subtitle_style,
        'h1': h1_style,
        'h2': h2_style,
        'body': body_style,
        'bullet': bullet_style,
        'callout': callout_style,
        'th': table_header,
        'td': table_cell,
        'td_bold': table_cell_bold
    }


def add_docx_header(doc, title, subtitle):
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.font.name = 'Arial'
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)
    p.paragraph_format.space_after = Pt(2)
    
    p2 = doc.add_paragraph()
    run2 = p2.add_run(subtitle)
    run2.font.name = 'Arial'
    run2.font.size = Pt(11)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
    p2.paragraph_format.space_after = Pt(14)

def add_docx_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.bold = True
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
    else:
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)

def add_docx_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    r_body = p.add_run(text)
    r_body.font.name = 'Arial'
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

def add_docx_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Arial'
        r_pre.font.size = Pt(9.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    r_body = p.add_run(text)
    r_body.font.name = 'Arial'
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

def style_docx_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Set header row shading
    for i, cell in enumerate(table.rows[0].cells):
        shading = parse_xml(r'<w:shd {} w:fill="4338CA"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(9)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    # Alternate row colors
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg_hex = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg_hex))
            cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                    
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

print("Business helpers module compiled successfully.")
