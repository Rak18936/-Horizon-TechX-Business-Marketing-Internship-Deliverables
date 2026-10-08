import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib import colors

from biz_helpers import (
    OUTPUT_DIR, PRIMARY, SECONDARY, DARK_TEXT, LIGHT_BG, ALT_BG, BORDER_CLR,
    NumberedCanvas, get_pdf_styles, add_docx_header, add_docx_heading,
    add_docx_p, add_docx_bullet, style_docx_table
)

def build_task_1():
    pdf_filename = os.path.join(OUTPUT_DIR, "Task_1_Startup_Market_Research_Report.pdf")
    docx_filename = os.path.join(OUTPUT_DIR, "Task_1_Startup_Market_Research_Report.docx")
    
    st = get_pdf_styles()
    story = []
    
    # Header
    story.append(Paragraph("TASK 1: STARTUP MARKET RESEARCH REPORT", st['title']))
    story.append(Paragraph("Horizon TechX — Industry Landscape, SWOT & Growth Opportunity Insights", st['subtitle']))
    story.append(HRFlowable(width="100%", thickness=1.5, color=PRIMARY, spaceAfter=12))
    
    # Executive Summary
    story.append(Paragraph("1. Executive Summary & Market Sizing", st['h1']))
    story.append(Paragraph(
        "The global EdTech and AI-driven skill acceleration market is experiencing an unprecedented paradigm shift. "
        "Driven by rapid technological advancements in Generative AI, cloud computing, and automated workflows, "
        "the global EdTech market is projected to grow from $142 Billion in 2024 to over $340 Billion by 2030 at a CAGR of 15.8%. "
        "This report provides a market research analysis for <b>Horizon TechX</b>, evaluating industry trends, competitive positioning, "
        "SWOT metrics, and high-impact growth opportunities.",
        st['body']
    ))
    story.append(Spacer(1, 6))
    
    market_table = [
        [Paragraph("Market Segment", st['th']), Paragraph("Valuation (USD)", st['th']), Paragraph("Growth Rate (CAGR)", st['th']), Paragraph("Horizon TechX Addressable Focus", st['th'])],
        [Paragraph("TAM (Total Addressable Market)", st['td_bold']), Paragraph("$340 Billion", st['td']), Paragraph("15.8%", st['td']), Paragraph("Global Online Learning & Tech Upskilling Market.", st['td'])],
        [Paragraph("SAM (Serviceable Addressable Market)", st['td_bold']), Paragraph("$24 Billion", st['td']), Paragraph("18.2%", st['td']), Paragraph("Tech Upskilling & Internship Programs in India & SE Asia.", st['td'])],
        [Paragraph("SOM (Serviceable Obtainable Market)", st['td_bold']), Paragraph("$450 Million", st['td']), Paragraph("28.4%", st['td']), Paragraph("Early Career Developers & Tech Interns (Years 1-3).", st['td'])]
    ]
    t_m = Table(market_table, colWidths=[140, 94, 110, 160])
    t_m.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE')
    ]))
    story.append(t_m)
    story.append(Spacer(1, 10))
    
    # SWOT Analysis
    story.append(Paragraph("2. Strategic SWOT Analysis", st['h1']))
    story.append(Paragraph(
        "To formulate a robust market penetration strategy, a SWOT analysis was conducted evaluating internal capabilities and external market dynamics:",
        st['body']
    ))
    
    swot_table = [
        [Paragraph("Strengths (Internal)", st['th']), Paragraph("Weaknesses (Internal)", st['th'])],
        [Paragraph("• Hands-on project-based internship model.<br/>• Industry-aligned curriculum (AI, HR, Data Science).<br/>• QR-verified credentials & placement support.<br/>• Highly affordable pricing compared to legacy bootcamps.", st['td']),
         Paragraph("• Emerging brand awareness in Tier-2/3 cities.<br/>• Reliance on third-party communication channels.<br/>• Limited enterprise B2B partner network.<br/>• Scalability limits on manual code review feedback.", st['td'])],
        [Paragraph("Opportunities (External)", st['th']), Paragraph("Threats (External)", st['th'])],
        [Paragraph("• Surge in demand for practical AI & Data Analytics skills.<br/>• Expansion into B2B corporate upskilling programs.<br/>• Strategic university & college campus partnerships.<br/>• Micro-credentialing & digital badge integration.", st['td']),
         Paragraph("• Intense competition from legacy platforms (Coursera, UpGrad).<br/>• Rapid shifts in tech stack frameworks.<br/>• Price wars from free YouTube/Open-Source content.<br/>• Economic slowdown impacting hiring volumes.", st['td'])]
    ]
    t_swot = Table(swot_table, colWidths=[252, 252])
    t_swot.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,0), PRIMARY),
        ('BACKGROUND', (1,0), (1,0), SECONDARY),
        ('BACKGROUND', (0,2), (0,2), SECONDARY),
        ('BACKGROUND', (1,2), (1,2), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('PADDING', (0,0), (-1,-1), 6),
        ('VALIGN', (0,0), (-1,-1), 'TOP')
    ]))
    story.append(t_swot)
    story.append(Spacer(1, 10))
    
    # Competitor Benchmark Matrix
    story.append(Paragraph("3. Competitor Analysis Matrix", st['h1']))
    story.append(Paragraph(
        "Horizon TechX competes against established MOOC platforms and specialized bootcamps. The matrix below benchmarks key operational parameters:",
        st['body']
    ))
    
    comp_table = [
        [Paragraph("Competitor", st['th']), Paragraph("Primary Product", st['th']), Paragraph("Pricing Point", st['th']), Paragraph("Placement Support", st['th']), Paragraph("Horizon TechX Advantage", st['th'])],
        [Paragraph("Coursera / edX", st['td_bold']), Paragraph("Video Courses", st['td']), Paragraph("₹3,500 – ₹8,000 / mo", st['td']), Paragraph("Self-serve job board", st['td']), Paragraph("Live internship projects & direct mentorship.", st['td'])],
        [Paragraph("UpGrad / Scaler", st['td_bold']), Paragraph("Long Bootcamps", st['td']), Paragraph("₹1.5L – ₹3.5L total", st['td']), Paragraph("Dedicated placement cell", st['td']), Paragraph("Fraction of cost with immediate practical exposure.", st['td'])],
        [Paragraph("Internshala", st['td_bold']), Paragraph("Internship Portal", st['td']), Paragraph("Freemium model", st['td']), Paragraph("Application listing", st['td']), Paragraph("Structured portfolio building & verified certification.", st['td'])],
        [Paragraph("Horizon TechX", st['td_bold']), Paragraph("Project Internships", st['td']), Paragraph("Accessible / Freemium", st['td']), Paragraph("Resume & Placement", st['td']), Paragraph("High ROI, project-first, fast-track skill validation.", st['td'])]
    ]
    t_comp = Table(comp_table, colWidths=[90, 95, 95, 100, 124])
    t_comp.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_CLR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, ALT_BG]),
        ('PADDING', (0,0), (-1,-1), 5),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE')
    ]))
    story.append(t_comp)
    story.append(Spacer(1, 10))
    
    # Strategic Growth Recommendations
    story.append(Paragraph("4. High-Growth Opportunity Insights & Action Plan", st['h1']))
    story.append(Paragraph("• <b>AI-Driven Personalization:</b> Integrate automated AI code reviewer bots to provide real-time feedback on intern project submissions, improving turnaround time by 80%.", st['bullet']))
    story.append(Paragraph("• <b>B2B Enterprise Partnerships:</b> Launch corporate upskilling tracks for IT services companies seeking fast-track onboarding for junior engineers.", st['bullet']))
    story.append(Paragraph("• <b>College Campus Ambassador Network:</b> Establish student ambassador chapters across 200+ engineering & management institutes to drive viral word-of-mouth acquisition.", st['bullet']))
    story.append(Paragraph("• <b>Verified Digital Badging:</b> Partner with LinkedIn and Credly to offer instant 1-click badge sharing for completed projects, amplifying organic social reach.", st['bullet']))
    
    # Build PDF
    doc_pdf = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        leftMargin=54, rightMargin=54,
        topMargin=54, bottomMargin=54
    )
    doc_pdf.build(story, canvasmaker=NumberedCanvas)
    print(f"Generated PDF: {pdf_filename}")

    # Build DOCX
    doc_x = docx.Document()
    for section in doc_x.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    add_docx_header(doc_x, "TASK 1: STARTUP MARKET RESEARCH REPORT", "Horizon TechX — Industry Landscape, SWOT & Growth Opportunity Insights")
    
    add_docx_heading(doc_x, "1. Executive Summary & Market Sizing", 1)
    add_docx_p(doc_x, "The global EdTech and AI-driven skill acceleration market is experiencing an unprecedented paradigm shift. Driven by rapid technological advancements in Generative AI, cloud computing, and automated workflows, the global EdTech market is projected to grow from $142 Billion in 2024 to over $340 Billion by 2030 at a CAGR of 15.8%. This report provides a market research analysis for Horizon TechX, evaluating industry trends, competitive positioning, SWOT metrics, and high-impact growth opportunities.")
    
    t_m_x = doc_x.add_table(rows=4, cols=4)
    t_m_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_matrix = [
        ["Market Segment", "Valuation (USD)", "Growth Rate (CAGR)", "Horizon TechX Addressable Focus"],
        ["TAM (Total Addressable Market)", "$340 Billion", "15.8%", "Global Online Learning & Tech Upskilling Market."],
        ["SAM (Serviceable Addressable Market)", "$24 Billion", "18.2%", "Tech Upskilling & Internship Programs in India & SE Asia."],
        ["SOM (Serviceable Obtainable Market)", "$450 Million", "28.4%", "Early Career Developers & Tech Interns (Years 1-3)."]
    ]
    for r_i, row in enumerate(m_matrix):
        for c_i, val in enumerate(row):
            t_m_x.cell(r_i, c_i).text = val
    style_docx_table(t_m_x)
    
    add_docx_heading(doc_x, "2. Strategic SWOT Analysis", 1)
    add_docx_p(doc_x, "To formulate a robust market penetration strategy, a SWOT analysis was conducted evaluating internal capabilities and external market dynamics:")
    
    s_matrix = [
        ["Strengths (Internal)", "Weaknesses (Internal)"],
        ["Hands-on project-based internship model. Industry-aligned curriculum (AI, HR, Data Science). QR-verified credentials & placement support. Highly affordable pricing.",
         "Emerging brand awareness in Tier-2/3 cities. Reliance on third-party channels. Limited enterprise B2B partner network. Manual code review scalability."],
        ["Opportunities (External)", "Threats (External)"],
        ["Surge in demand for practical AI & Data Analytics skills. Expansion into B2B corporate upskilling programs. Campus ambassador network across colleges.",
         "Intense competition from legacy platforms. Rapid shifts in tech stack frameworks. Price wars from free YouTube/Open-Source content."]
    ]
    t_swot_x = doc_x.add_table(rows=len(s_matrix), cols=2)
    t_swot_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_i, row in enumerate(s_matrix):
        for c_i, val in enumerate(row):
            t_swot_x.cell(r_i, c_i).text = val
    style_docx_table(t_swot_x)
    
    add_docx_heading(doc_x, "3. Competitor Analysis Matrix", 1)
    add_docx_p(doc_x, "Horizon TechX competes against established MOOC platforms and specialized bootcamps. The matrix below benchmarks key operational parameters:")
    
    t_c_x = doc_x.add_table(rows=5, cols=5)
    t_c_x.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_matrix = [
        ["Competitor", "Primary Product", "Pricing Point", "Placement Support", "Horizon TechX Advantage"],
        ["Coursera / edX", "Video Courses", "₹3,500 – ₹8,000 / mo", "Self-serve job board", "Live internship projects & direct mentorship."],
        ["UpGrad / Scaler", "Long Bootcamps", "₹1.5L – ₹3.5L total", "Dedicated placement cell", "Fraction of cost with immediate practical exposure."],
        ["Internshala", "Internship Portal", "Freemium model", "Application listing", "Structured portfolio building & verified certification."],
        ["Horizon TechX", "Project Internships", "Accessible / Freemium", "Resume & Placement", "High ROI, project-first, fast-track skill validation."]
    ]
    for r_i, row in enumerate(c_matrix):
        for c_i, val in enumerate(row):
            t_c_x.cell(r_i, c_i).text = val
    style_docx_table(t_c_x)
    
    add_docx_heading(doc_x, "4. High-Growth Opportunity Insights & Action Plan", 1)
    add_docx_bullet(doc_x, " Integrate automated AI code reviewer bots to provide real-time feedback on intern project submissions.", "AI-Driven Personalization:")
    add_docx_bullet(doc_x, " Launch corporate upskilling tracks for IT services companies seeking fast-track onboarding for junior engineers.", "B2B Enterprise Partnerships:")
    add_docx_bullet(doc_x, " Establish student ambassador chapters across 200+ engineering & management institutes.", "College Campus Ambassador Network:")
    add_docx_bullet(doc_x, " Partner with LinkedIn and Credly to offer instant 1-click badge sharing for completed projects.", "Verified Digital Badging:")
    
    doc_x.save(docx_filename)
    print(f"Generated DOCX: {docx_filename}")

if __name__ == "__main__":
    build_task_1()
